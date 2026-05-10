from flask import Flask, render_template, redirect, url_for, request, flash, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from werkzeug.security import generate_password_hash, check_password_hash
import os
import io
from datetime import datetime, timedelta
import pytz
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
try:
    import pandas as pd
except ImportError:
    pd = None

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'garden-of-eden-secret-key')
app.config['SQLALCHEMY_DATABASE_URI'] = os.environ.get('DATABASE_URL', 'sqlite:///adenticket.db')

# Fix for Render/Heroku postgres:// vs postgresql://
if app.config['SQLALCHEMY_DATABASE_URI']:
    if app.config['SQLALCHEMY_DATABASE_URI'].startswith("postgres://"):
        app.config['SQLALCHEMY_DATABASE_URI'] = app.config['SQLALCHEMY_DATABASE_URI'].replace("postgres://", "postgresql://", 1)
    
    # Add SSL requirement for Supabase/External DBs if not present
    if "supabase.co" in app.config['SQLALCHEMY_DATABASE_URI'] and "sslmode" not in app.config['SQLALCHEMY_DATABASE_URI']:
        if "?" in app.config['SQLALCHEMY_DATABASE_URI']:
            app.config['SQLALCHEMY_DATABASE_URI'] += "&sslmode=require"
        else:
            app.config['SQLALCHEMY_DATABASE_URI'] += "?sslmode=require"

app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)
login_manager = LoginManager()
login_manager.login_view = 'login'
login_manager.init_app(app)

# Models
class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(150), unique=True, nullable=False)
    password = db.Column(db.String(150), nullable=False)
    is_admin = db.Column(db.Boolean, default=False)
    role = db.Column(db.String(50), default='user')
    is_banned = db.Column(db.Boolean, default=False)
    profile_image = db.Column(db.String(500), default='images/profile_default.jpg')
    nickname = db.Column(db.String(150), nullable=True)
    status_text = db.Column(db.String(500), nullable=True)
    submissions = db.relationship('TicketForm', backref='user', lazy=True)

    @property
    def is_actually_admin(self):
        # Always allow 'admin' username to access admin panel regardless of role
        is_admin_bool = self.is_admin in [True, 1, '1', 'true', 'True']
        return is_admin_bool or self.role in ['admin', 'developer', 'admin_miru', 'admin_focus', 'admin_assistant'] or self.username == 'admin'

class Concert(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False)
    price = db.Column(db.String(500), nullable=False)
    image = db.Column(db.String(500), nullable=False)
    max_capacity = db.Column(db.Integer, default=0)
    current_bookings = db.Column(db.Integer, default=0)
    status_badge = db.Column(db.String(50), nullable=True, default='none')
    is_open = db.Column(db.Boolean, default=True)
    date = db.Column(db.String(200), nullable=True, default='')
    location = db.Column(db.String(500), nullable=True, default='')

class TicketForm(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    customer_name = db.Column(db.String(500), nullable=False)
    customer_surname = db.Column(db.String(500), nullable=False)
    email = db.Column(db.String(500), nullable=False)
    phone_number = db.Column(db.String(500), nullable=False)
    twitter_account = db.Column(db.String(500), nullable=False)
    membership = db.Column(db.String(500), nullable=True)
    show_time = db.Column(db.String(500), nullable=False)
    ticket_name = db.Column(db.String(500), nullable=True)
    ticket_surname = db.Column(db.String(500), nullable=True)
    payment_method = db.Column(db.String(500), nullable=False)
    ticket_price = db.Column(db.String(500), nullable=False)
    backup_price = db.Column(db.String(500), nullable=True)
    primary_zone = db.Column(db.String(500), nullable=False)
    backup_zone = db.Column(db.String(500), nullable=True)
    concert_name = db.Column(db.String(200), nullable=False)
    ticket_quantity = db.Column(db.String(500), default="1")
    status = db.Column(db.String(20), default='Pending')
    note = db.Column(db.String(1000), nullable=True)
    ticketing_platform = db.Column(db.String(500), nullable=True)
    created_at = db.Column(db.DateTime, default=lambda: datetime.utcnow() + timedelta(hours=7))

class SystemSetting(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    key = db.Column(db.String(50), unique=True, nullable=False)
    value = db.Column(db.String(500), nullable=False)

@login_manager.user_loader
def load_user(user_id):
    try:
        return User.query.get(int(user_id))
    except Exception:
        return None

# Routes
@app.before_request
def check_maintenance():
    # Allow access to static files, login, logout, and migrate
    if request.path.startswith('/static') or \
       request.path in [url_for('login'), url_for('logout'), url_for('migrate'), '/maintenance']:
        return

    # Check maintenance mode
    maintenance = SystemSetting.query.filter_by(key='maintenance_mode').first()
    if maintenance and maintenance.value == 'True':
        # If user is not admin, redirect to maintenance page
        if not current_user.is_authenticated or not current_user.is_actually_admin:
            return render_template('maintenance.html')

@app.route('/maintenance')
def maintenance_page():
    return render_template('maintenance.html')

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/concerts')
def concerts_page():
    concerts = Concert.query.all()
    booked_concerts = []
    if current_user.is_authenticated:
        booked_concerts = [s.concert_name for s in TicketForm.query.filter_by(user_id=current_user.id).all()]
    return render_template('concerts.html', concerts=concerts, booked_concerts=booked_concerts)

@app.context_processor
def inject_settings():
    announcement = SystemSetting.query.filter_by(key='announcement').first()
    return dict(announcement=announcement.value if announcement else "")

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()
        user = User.query.filter_by(username=username).first()
        if user and check_password_hash(user.password, password):
            if user.is_banned:
                flash('Your account has been banned.', 'error')
                return redirect(url_for('login'))
            login_user(user)
            return redirect(url_for('index'))
        else:
            flash('Login failed. Check your username and password.')
    return render_template('login.html')

@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()
        hashed_password = generate_password_hash(password, method='pbkdf2:sha256')
        new_user = User(username=username, password=hashed_password, is_admin=False)
        try:
            db.session.add(new_user)
            db.session.commit()
            flash('Registration successful!')
            return redirect(url_for('login'))
        except Exception as e:
            db.session.rollback()
            if 'UNIQUE constraint failed' in str(e) or 'unique constraint' in str(e).lower():
                flash('ชื่อผู้ใช้นี้ถูกใช้งานแล้ว กรุณาเลือกชื่ออื่น', 'error')
            else:
                flash(f'Error: {str(e)}', 'error')
    return render_template('register.html')

@app.route('/profile/update', methods=['POST'])
@login_required
def update_profile():
    nickname = request.form.get('nickname')
    status_text = request.form.get('status_text')
    
    image_file = request.files.get('profile_image')
    if image_file and image_file.filename:
        upload_folder = os.path.join('static', 'uploads', 'profiles')
        if not os.path.exists(upload_folder):
            os.makedirs(upload_folder)
        filename = f"user_{current_user.id}_{image_file.filename}"
        image_path = os.path.join(upload_folder, filename)
        image_file.save(image_path)
        current_user.profile_image = f"uploads/profiles/{filename}"
        
    current_user.nickname = nickname
    current_user.status_text = status_text
    db.session.commit()
    flash('อัปเดตโปรไฟล์เรียบร้อยแล้ว!', 'success_modal')
    return redirect(request.referrer or url_for('index'))

@app.route('/reset-password', methods=['GET', 'POST'])
def reset_password():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        new_password = request.form.get('new_password', '').strip()
        confirm_password = request.form.get('confirm_password', '').strip()

        if not username or not new_password:
            flash('กรุณากรอกข้อมูลให้ครบถ้วน', 'reset_error')
            return redirect(url_for('reset_password'))

        if new_password != confirm_password:
            flash('รหัสผ่านใหม่ไม่ตรงกัน', 'reset_error')
            return redirect(url_for('reset_password'))

        if len(new_password) < 4:
            flash('รหัสผ่านต้องมีอย่างน้อย 4 ตัวอักษร', 'reset_error')
            return redirect(url_for('reset_password'))

        user = User.query.filter_by(username=username).first()
        if not user:
            flash('ไม่พบชื่อผู้ใช้นี้ในระบบ', 'reset_error')
            return redirect(url_for('reset_password'))

        user.password = generate_password_hash(new_password, method='pbkdf2:sha256')
        db.session.commit()
        flash('เปลี่ยนรหัสผ่านเรียบร้อยแล้ว! กรุณาเข้าสู่ระบบด้วยรหัสใหม่', 'reset_success')
        return redirect(url_for('reset_password'))
        
    return render_template('reset.html')


@app.route('/submit_form', methods=['POST'])
@login_required
def submit_form():
    customer_name = request.form.get('first_name') # Map first_name to customer_name
    customer_surname = request.form.get('surname') # Map surname to customer_surname
    email = request.form.get('email')
    phone = request.form.get('phone')
    twitter = request.form.get('twitter')
    membership = request.form.get('membership')
    payment_method = request.form.get('payment_method')
    concert_name = request.form.get('concert_name')
    concert = Concert.query.filter_by(name=concert_name).first()

    # Ticketing Platform
    ticketing_platform = request.form.get('ticketing_platform')
    if ticketing_platform == 'Other':
        ticketing_platform = request.form.get('other_platform', 'Other')

    # Handle multiple show times
    show_times = request.form.getlist('show_times[]')
    show_time = "\n".join([t.strip() for t in show_times if t.strip()])
    if not show_time:
        show_time = request.form.get('show_time', '-')
    
    # Handle multiple ticket names/surnames
    ticket_names = request.form.getlist('ticket_names[]')
    ticket_surnames = request.form.getlist('ticket_surnames[]')
    
    # Combine each name with its corresponding surname
    full_name_list = []
    for n, s in zip(ticket_names, ticket_surnames):
        if n.strip() or s.strip():
            full_name_list.append(f"{n.strip()} {s.strip()}".strip())
    
    ticket_name = "\n".join(full_name_list)
    ticket_surname = "" # No longer needed separately
    
    # Handle multiple zones, prices and quantities
    primary_zones = request.form.getlist('primary_zones[]')
    backup_zones = request.form.getlist('backup_zones[]')
    ticket_prices = request.form.getlist('ticket_prices[]')
    backup_prices = request.form.getlist('backup_prices[]')
    ticket_quantities = request.form.getlist('ticket_quantities[]')
    
    primary_zone = "\n".join([z.strip() for z in primary_zones if z.strip()])
    backup_zone = "\n".join([z.strip() for z in backup_zones if z.strip()])
    ticket_price = "\n".join([p.strip() for p in ticket_prices if p.strip()])
    backup_price = "\n".join([p.strip() for p in backup_prices if p.strip()])
    
    # Calculate total quantity for capacity check, but store as multi-line string
    total_qty = 0
    qty_list = []
    for q in ticket_quantities:
        if q.strip():
            try:
                val = int(q)
                total_qty += val
                qty_list.append(str(val))
            except:
                qty_list.append("1")
                total_qty += 1
    
    ticket_quantity_str = "\n".join(qty_list)
    note = request.form.get('note', '')
    
    # Check for duplicate submission (Allow up to 2 per concert per user)
    existing_submissions_count = TicketForm.query.filter_by(user_id=current_user.id, concert_name=concert_name).count()
    if existing_submissions_count >= 2:
        flash('คุณได้ลงทะเบียนงานนี้ครบ 2 รอบแล้ว หากต้องการแก้ไขข้อมูล กรุณาไปที่เมนู "รายการจองของฉัน"', 'error')
        return redirect(url_for('my_bookings'))

    if concert and not concert.is_open:
        flash('ขออภัย งานนี้ปิดรับจองแล้ว', 'error')
        return redirect(url_for('index'))

    if concert and concert.max_capacity > 0:
        if (concert.current_bookings + total_qty) > concert.max_capacity:
            flash(f'Sorry, this concert only has {concert.max_capacity - concert.current_bookings} seats left.', 'error')
            return redirect(url_for('index'))
            
    new_ticket = TicketForm(
        user_id=current_user.id,
        customer_name=customer_name,
        customer_surname=customer_surname,
        email=email,
        phone_number=phone,
        twitter_account=twitter,
        membership=membership,
        show_time=show_time,
        ticket_name=ticket_name,
        ticket_surname=ticket_surname,
        payment_method=payment_method,
        ticket_price=ticket_price,
        backup_price=backup_price if backup_price else None,
        primary_zone=primary_zone,
        backup_zone=backup_zone,
        concert_name=concert_name,
        ticket_quantity=ticket_quantity_str,
        note=note,
        ticketing_platform=ticketing_platform
    )
    
    if concert:
        concert.current_bookings += total_qty
        
    db.session.add(new_ticket)
    db.session.commit()
    flash('booking_success', 'success_modal')
    return redirect(url_for('index'))

@app.route('/admin')
@login_required
def admin():
    if not current_user.is_actually_admin:
        flash('Unauthorized access.')
        return redirect(url_for('index'))
    concerts = Concert.query.all()
    maintenance = SystemSetting.query.filter_by(key='maintenance_mode').first()
    maintenance_status = maintenance.value == 'True' if maintenance else False
    return render_template('admin.html', concerts=concerts, maintenance_status=maintenance_status)

@app.route('/admin/submissions')
@login_required
def admin_submissions():
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    submissions = TicketForm.query.order_by(TicketForm.concert_name, TicketForm.id.desc()).all()
    return render_template('admin_submissions.html', submissions=submissions)

@app.route('/admin/concert/add', methods=['POST'])
@login_required
def add_concert():
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    
    name = request.form.get('name')
    price = request.form.get('price')
    image_url = request.form.get('image_url')
    max_capacity = request.form.get('max_capacity', 0)
    status_badge = request.form.get('status_badge', 'none')
    
    image_file = request.files.get('image_file')
    if image_file and image_file.filename:
        upload_folder = os.path.join('static', 'uploads')
        if not os.path.exists(upload_folder):
            os.makedirs(upload_folder)
        filename = image_file.filename
        image_path = os.path.join(upload_folder, filename)
        image_file.save(image_path)
        image_url = url_for('static', filename='uploads/' + filename)

    date = request.form.get('date', '')
    location = request.form.get('location', '')

    new_concert = Concert(
        name=name, 
        price=price, 
        image=image_url, 
        max_capacity=int(max_capacity), 
        status_badge=status_badge,
        date=date,
        location=location
    )
    db.session.add(new_concert)
    db.session.commit()
    flash('Concert added successfully!')
    return redirect(url_for('admin'))

@app.route('/admin/toggle_maintenance', methods=['POST'])
@login_required
def toggle_maintenance():
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    
    maintenance = SystemSetting.query.filter_by(key='maintenance_mode').first()
    if not maintenance:
        maintenance = SystemSetting(key='maintenance_mode', value='False')
        db.session.add(maintenance)
    
    if maintenance.value == 'True':
        maintenance.value = 'False'
        flash('Maintenance mode disabled!')
    else:
        maintenance.value = 'True'
        flash('Maintenance mode enabled!')
    
    db.session.commit()
    return redirect(url_for('admin'))

@app.route('/admin/update_announcement', methods=['POST'])
@login_required
def update_announcement():
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    new_text = request.form.get('announcement_text', '').strip()
    announcement = SystemSetting.query.filter_by(key='announcement').first()
    if not announcement:
        announcement = SystemSetting(key='announcement', value='')
        db.session.add(announcement)
    announcement.value = new_text
    db.session.commit()
    flash('Announcement updated!')
    return redirect(url_for('admin'))

@app.route('/admin/delete_announcement', methods=['POST'])
@login_required
def delete_announcement():
    if not current_user.is_actually_admin:
        flash('Access Denied', 'error')
        return redirect(url_for('index'))
    SystemSetting.query.filter_by(key='announcement').delete()
    db.session.commit()
    flash('ลบประกาศเรียบร้อยแล้ว')
    return redirect(url_for('admin'))

@app.route('/admin/concert/delete/<int:id>', methods=['POST'])
@login_required
def delete_concert(id):
    if not current_user.is_actually_admin:
        flash('Access Denied', 'error')
        return redirect(url_for('index'))
    concert = Concert.query.get_or_404(id)
    db.session.delete(concert)
    db.session.commit()
    flash('ลบคอนเสิร์ตเรียบร้อยแล้ว')
    return redirect(url_for('admin'))

@app.route('/admin/concert/toggle_booking/<int:id>')
@login_required
def toggle_concert_booking(id):
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    concert = Concert.query.get_or_404(id)
    concert.is_open = not concert.is_open
    db.session.commit()
    status = "เปิดจองแล้ว" if concert.is_open else "ปิดจองแล้ว"
    flash(f'งาน "{concert.name}" {status}')
    return redirect(url_for('admin'))

@app.route('/admin/concert/edit/<int:id>', methods=['POST'])
@login_required
def edit_concert(id):
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    concert = Concert.query.get_or_404(id)
    concert.name = request.form.get('name', concert.name)
    concert.price = request.form.get('price', concert.price)
    new_image = request.form.get('image_url', '').strip()
    if new_image:
        concert.image = new_image
    concert.max_capacity = int(request.form.get('max_capacity', concert.max_capacity))
    concert.status_badge = request.form.get('status_badge', concert.status_badge)
    concert.date = request.form.get('date', concert.date)
    concert.location = request.form.get('location', concert.location)
    if 'is_open' in request.form:
        concert.is_open = request.form.get('is_open') == 'True'
    db.session.commit()
    flash(f'Updated "{concert.name}" successfully!')
    return redirect(url_for('admin'))

@app.route('/admin/submissions/delete/<int:id>', methods=['POST'])
@login_required
def delete_submission(id):
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    submission = TicketForm.query.get_or_404(id)
    concert = Concert.query.filter_by(name=submission.concert_name).first()
    if concert:
        qty = sum(int(q) for q in submission.ticket_quantity.split('\n') if q.strip())
        concert.current_bookings = max(0, concert.current_bookings - qty)
    db.session.delete(submission)
    db.session.commit()
    flash('ลบข้อมูลรายการจองเรียบร้อยแล้ว', 'success_modal')
    return redirect(url_for('admin_submissions'))

@app.route('/admin/submissions/delete_all', methods=['POST'])
@login_required
def delete_all_submissions():
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    TicketForm.query.delete()
    concerts = Concert.query.all()
    for c in concerts:
        c.current_bookings = 0
    db.session.commit()
    flash('ลบข้อมูลรายการจองทั้งหมดเรียบร้อยแล้ว', 'success_modal')
    return redirect(url_for('admin_submissions'))

@app.route('/admin/submissions/edit/<int:id>', methods=['POST'])
@login_required
def edit_submission(id):
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    submission = TicketForm.query.get_or_404(id)
    
    # Update fields only if they exist in the form
    if 'customer_name' in request.form: submission.customer_name = request.form['customer_name']
    if 'customer_surname' in request.form: submission.customer_surname = request.form['customer_surname']
    if 'email' in request.form: submission.email = request.form['email']
    if 'phone_number' in request.form: submission.phone_number = request.form['phone_number']
    if 'twitter_account' in request.form: submission.twitter_account = request.form['twitter_account']
    if 'membership' in request.form: submission.membership = request.form['membership']
    if 'show_time' in request.form: submission.show_time = request.form['show_time']
    if 'ticket_name' in request.form: submission.ticket_name = request.form['ticket_name']
    if 'ticket_surname' in request.form: submission.ticket_surname = request.form['ticket_surname']
    if 'payment_method' in request.form: submission.payment_method = request.form['payment_method']
    if 'primary_zone' in request.form: submission.primary_zone = request.form['primary_zone']
    if 'backup_zone' in request.form: submission.backup_zone = request.form['backup_zone']
    if 'ticket_quantity' in request.form: submission.ticket_quantity = request.form['ticket_quantity']
    if 'note' in request.form:
        submission.note = request.form['note']
    if 'ticketing_platform' in request.form:
        submission.ticketing_platform = request.form['ticketing_platform']
    
    if request.form.get('ticket_price'):
        submission.ticket_price = request.form.get('ticket_price')
    if request.form.get('backup_price'):
        submission.backup_price = request.form.get('backup_price')
    elif 'backup_price' in request.form:
        submission.backup_price = None

    new_status = request.form.get('status')
    if new_status:
        submission.status = new_status
        
    db.session.commit()
    flash('อัปเดตข้อมูลรายการจองเรียบร้อยแล้ว', 'success_modal')
    return redirect(url_for('admin_submissions'))

@app.route('/admin/export/submissions')
@login_required
def export_submissions():
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    
    if pd is None:
        flash('Error: pandas and openpyxl are not installed. Please run "pip install pandas openpyxl" on the server.', 'error')
        return redirect(url_for('admin_submissions'))

    submissions = TicketForm.query.order_by(TicketForm.id.desc()).all()
    
    data = []
    for s in submissions:
        data.append({
            'Order ID': s.id,
            'Platform (เว็บที่กด)': s.ticketing_platform,
            'User / ID (ลูกค้า)': s.customer_name,
            'Password (ลูกค้า)': s.customer_surname,
            'Concert Name': s.concert_name,
            'Show Time (รอบ)': s.show_time,
            'Quantity (จำนวน)': s.ticket_quantity,
            'Primary Zone': s.primary_zone,
            'Backup Zone': s.backup_zone,
            'Price (ราคาหลัก)': s.ticket_price,
            'Backup Price (ราคาสำรอง)': s.backup_price,
            'Ticket Name (ชื่อบนบัตร)': s.ticket_name,
            'Customer Phone': s.phone_number,
            'Twitter / X': s.twitter_account,
            'Membership Info': s.membership,
            'Email': s.email,
            'Payment Method': s.payment_method,
            'Booking Status': s.status,
            'Submission Time': s.created_at.strftime('%Y-%m-%d %H:%M:%S') if s.created_at else '-',
            'Note (หมายเหตุ)': s.note
        })
    
    df = pd.DataFrame(data)
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Submissions')
        
        # Access the openpyxl workbook and sheet to style it
        workbook = writer.book
        worksheet = writer.sheets['Submissions']
        
        # Styles
        header_fill = PatternFill(start_color="FF758C", end_color="FF758C", fill_type="solid")
        header_font = Font(name='Calibri', color="FFFFFF", bold=True, size=12)
        content_font = Font(name='Calibri', size=11)
        center_alignment = Alignment(horizontal="center", vertical="center")
        left_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        
        # Style Header
        worksheet.row_dimensions[1].height = 25
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment
            cell.border = thin_border
            
        # Style Content and Auto-Adjust Widths
        for i, col in enumerate(df.columns, 1):
            max_len = len(str(col)) + 4
            for row_idx in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row_idx, column=i)
                cell.font = content_font
                cell.alignment = left_alignment
                cell.border = thin_border
                
                # Auto-calculate width
                val = str(cell.value) if cell.value is not None else ""
                lines = val.split('\n')
                line_max = max([len(l) for l in lines]) if lines else 0
                max_len = max(max_len, line_max + 2)
            
            # Apply width with a reasonable limit
            col_letter = get_column_letter(i)
            worksheet.column_dimensions[col_letter].width = min(max_len, 45)
        

    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='adenticket_submissions.xlsx'
    )

@app.route('/my-bookings/edit/<int:id>', methods=['POST'])
@login_required
def update_booking(id):
    submission = TicketForm.query.get_or_404(id)
    if submission.user_id != current_user.id and not current_user.is_actually_admin:
        flash('Unauthorized access.', 'error')
        return redirect(url_for('my_bookings'))
    
    # Only allow editing if status is Pending (optional but recommended)
    if submission.status != 'Pending' and not current_user.is_actually_admin:
        flash('Cannot edit a confirmed or cancelled booking.', 'error')
        return redirect(url_for('my_bookings'))

    if 'customer_name' in request.form: submission.customer_name = request.form['customer_name']
    if 'customer_surname' in request.form: submission.customer_surname = request.form['customer_surname']
    if 'email' in request.form: submission.email = request.form['email']
    if 'phone_number' in request.form: submission.phone_number = request.form['phone_number']
    if 'twitter_account' in request.form: submission.twitter_account = request.form['twitter_account']
    if 'membership' in request.form: submission.membership = request.form['membership']
    if 'show_time' in request.form: submission.show_time = request.form['show_time']
    if 'ticket_name' in request.form: submission.ticket_name = request.form['ticket_name']
    if 'ticket_surname' in request.form: submission.ticket_surname = request.form['ticket_surname']
    if 'payment_method' in request.form: submission.payment_method = request.form['payment_method']
    if 'primary_zone' in request.form: submission.primary_zone = request.form['primary_zone']
    if 'backup_zone' in request.form: submission.backup_zone = request.form['backup_zone']
    
    if 'ticket_quantity' in request.form:
        try:
            new_qty = int(request.form['ticket_quantity'])
            old_qty = submission.ticket_quantity or 1
            diff = new_qty - old_qty
            concert = Concert.query.filter_by(name=submission.concert_name).first()
            if concert:
                if diff > 0 and concert.max_capacity > 0:
                    if (concert.current_bookings + diff) > concert.max_capacity:
                        flash(f'Cannot increase quantity. Only {concert.max_capacity - concert.current_bookings} seats left.', 'error')
                        return redirect(url_for('my_bookings'))
                concert.current_bookings += diff
            submission.ticket_quantity = new_qty
        except: pass

    if 'note' in request.form:
        submission.note = request.form['note']
    if 'ticketing_platform' in request.form:
        submission.ticketing_platform = request.form['ticketing_platform']

    if request.form.get('ticket_price'):
        submission.ticket_price = float(request.form.get('ticket_price'))
    if request.form.get('backup_price'):
        submission.backup_price = float(request.form.get('backup_price'))
    elif 'backup_price' in request.form:
        submission.backup_price = None
    
    db.session.commit()
    flash('ข้อมูลการจองของคุณได้รับการแก้ไขแล้ว', 'success_modal')
    return redirect(url_for('my_bookings'))

@app.route('/admin/users')
@login_required
def admin_users():
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    users = User.query.all()
    return render_template('admin_users.html', users=users)

@app.route('/admin/users/role/<int:id>', methods=['POST'])
@login_required
def set_user_role(id):
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    user = User.query.get_or_404(id)
    # Allow self-promotion to developer if no developer exists yet (bootstrap case)
    no_developer_exists = User.query.filter_by(role='developer').count() == 0
    if user.id == current_user.id and current_user.role != 'developer' and not no_developer_exists:
        flash("Only developers can change their own rank or other admins!", "error")
        return redirect(url_for('admin_users'))
    
    new_role = request.form.get('role')
    allowed_roles = ['developer', 'admin_miru', 'admin_focus', 'admin_assistant', 'member', 'user', 'admin']
    if new_role in allowed_roles:
        user.role = new_role
        # Sync is_admin for compatibility
        user.is_admin = (new_role in ['admin', 'developer', 'admin_miru', 'admin_focus', 'admin_assistant'])
        db.session.commit()
        flash(f"Updated role for {user.username} to {new_role}")
    return redirect(url_for('admin_users'))

@app.route('/admin/users/promote/<int:id>')
@login_required
def promote_user(id):
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    user = User.query.get_or_404(id)
    if user.id == current_user.id:
        flash("You cannot change your own rank!", "error")
        return redirect(url_for('admin_users'))
    
    # Toggle between admin and user for the old button compatibility
    if user.role == 'admin':
        user.role = 'user'
        user.is_admin = False
    else:
        user.role = 'admin'
        user.is_admin = True
        
    db.session.commit()
    flash(f"Updated rank for {user.username}")
    return redirect(url_for('admin_users'))

@app.route('/admin/users/ban/<int:id>')
@login_required
def ban_user(id):
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    user = User.query.get_or_404(id)
    if user.id == current_user.id:
        flash("You cannot ban yourself!", "error")
        return redirect(url_for('admin_users'))
    user.is_banned = not user.is_banned
    db.session.commit()
    status = "banned" if user.is_banned else "unbanned"
    flash(f"User {user.username} has been {status}")
    return redirect(url_for('admin_users'))

@app.route('/admin/users/delete/<int:id>', methods=['POST'])
@login_required
def delete_user(id):
    if not current_user.is_actually_admin:
        return redirect(url_for('index'))
    user = User.query.get_or_404(id)
    if user.id == current_user.id:
        flash("You cannot delete yourself!", "error")
        return redirect(url_for('admin_users'))
    TicketForm.query.filter_by(user_id=user.id).delete()
    db.session.delete(user)
    db.session.commit()
    flash(f"User {user.username} and all their submissions have been deleted.")
    return redirect(url_for('admin_users'))

@app.route('/migrate')
def migrate():
    from sqlalchemy import text
    results = []
    
    # Columns to try and add (table, column, definition)
    columns = [
        ('user', 'is_banned', "BOOLEAN DEFAULT FALSE"),
        ('user', 'role', "VARCHAR(50) DEFAULT 'user'"),
        ('user', 'profile_image', "VARCHAR(500) DEFAULT 'images/profile_default.jpg'"),
        ('user', 'nickname', "VARCHAR(150)"),
        ('user', 'status_text', "VARCHAR(500)"),
        ('ticket_form', 'status', "VARCHAR(20) DEFAULT 'Pending'"),
        ('ticket_form', 'note', "VARCHAR(1000)"),
        ('ticket_form', 'ticketing_platform', "VARCHAR(500)"),
        ('concert', 'status_badge', "VARCHAR(50) DEFAULT 'none'"),
        ('concert', 'is_open', "BOOLEAN DEFAULT TRUE"),
        ('concert', 'date', "VARCHAR(200) DEFAULT ''"),
        ('concert', 'location', "VARCHAR(500) DEFAULT ''"),
        ('ticket_form', 'created_at', "TIMESTAMP"),
        ('ticket_form', 'ticket_quantity', "VARCHAR(500) DEFAULT '1'")
    ]

    for table, col, defn in columns:
        try:
            db.session.execute(text(f'ALTER TABLE "{table}" ADD COLUMN {col} {defn}'))
            db.session.commit()
            results.append(f"SUCCESS: Added {col} to {table}")
        except Exception as e:
            db.session.rollback()
            # If the error contains 'already exists' or 'duplicate column', it's safe to ignore
            err_msg = str(e).lower()
            if 'already exists' in err_msg or 'duplicate column' in err_msg or 'duplicate column name' in err_msg:
                results.append(f"OK: {col} in {table} (already exists)")
            else:
                results.append(f"FAIL: {col} in {table} - {str(e)}")

    # Specific fix for extending column lengths
    try:
        # PostgreSQL syntax for modifying column length
        db.session.execute(text('ALTER TABLE ticket_form ALTER COLUMN ticketing_platform TYPE VARCHAR(500)'))
        db.session.execute(text('ALTER TABLE ticket_form ALTER COLUMN phone_number TYPE VARCHAR(500)'))
        db.session.execute(text('ALTER TABLE ticket_form ALTER COLUMN twitter_account TYPE VARCHAR(500)'))
        db.session.execute(text('ALTER TABLE ticket_form ALTER COLUMN customer_name TYPE VARCHAR(500)'))
        db.session.execute(text('ALTER TABLE ticket_form ALTER COLUMN customer_surname TYPE VARCHAR(500)'))
        db.session.execute(text('ALTER TABLE ticket_form ALTER COLUMN email TYPE VARCHAR(500)'))
        db.session.execute(text('ALTER TABLE ticket_form ALTER COLUMN membership TYPE VARCHAR(500)'))
        db.session.execute(text('ALTER TABLE ticket_form ALTER COLUMN payment_method TYPE VARCHAR(500)'))
        db.session.commit()
        results.append(f"SUCCESS: Extended column lengths to 500 characters")
    except Exception as e:
        db.session.rollback()
        results.append(f"ALTER COL ERROR: {str(e)}")

    # Initialize Settings
    try:
        for key, val in [('maintenance_mode', 'False'), ('announcement', '')]:
            if not SystemSetting.query.filter_by(key=key).first():
                db.session.add(SystemSetting(key=key, value=val))
                db.session.commit()
                results.append(f"INITIALIZED: {key}")
    except Exception as e:
        results.append(f"ERROR SETTINGS: {e}")

    return "<h3>Migration Done</h3><ul><li>" + "</li><li>".join(results) + "</li></ul>"

@app.route('/my-bookings')
@login_required
def my_bookings():
    submissions = TicketForm.query.filter_by(user_id=current_user.id).all()
    return render_template('my_bookings.html', submissions=submissions)

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('index'))

with app.app_context():
    try:
        db.create_all()
        print("Database initialized successfully!")
        # Auto-migrate: add missing columns if they don't exist
        from sqlalchemy import text, inspect
        inspector = inspect(db.engine)
        columns_to_add = [
            ('user', 'is_banned', "BOOLEAN DEFAULT FALSE"),
            ('user', 'role', "VARCHAR(50) DEFAULT 'user'"),
            ('user', 'profile_image', "VARCHAR(500) DEFAULT 'images/profile_default.jpg'"),
            ('user', 'nickname', "VARCHAR(150)"),
            ('user', 'status_text', "VARCHAR(500)"),
            ('ticket_form', 'status', "VARCHAR(20) DEFAULT 'Pending'"),
            ('ticket_form', 'note', "VARCHAR(1000)"),
            ('ticket_form', 'ticketing_platform', "VARCHAR(500)"),
            ('concert', 'status_badge', "VARCHAR(50) DEFAULT 'none'"),
            ('concert', 'is_open', "BOOLEAN DEFAULT TRUE"),
            ('concert', 'date', "VARCHAR(200) DEFAULT ''"),
            ('concert', 'location', "VARCHAR(500) DEFAULT ''"),
            ('ticket_form', 'created_at', "TIMESTAMP"),
            ('ticket_form', 'ticket_quantity', "VARCHAR(500) DEFAULT '1'")
        ]
        for table, col, defn in columns_to_add:
            try:
                existing_cols = [c['name'] for c in inspector.get_columns(table)]
                if col not in existing_cols:
                    db.session.execute(text(f'ALTER TABLE "{table}" ADD COLUMN {col} {defn}'))
                    db.session.commit()
                    print(f"Auto-migrated: Added {col} to {table}")
            except Exception as me:
                db.session.rollback()
                print(f"Auto-migrate skip {col} in {table}: {me}")
        # Initialize default settings
        for key, val in [('maintenance_mode', 'False'), ('announcement', '')]:
            try:
                if not SystemSetting.query.filter_by(key=key).first():
                    db.session.add(SystemSetting(key=key, value=val))
                    db.session.commit()
            except Exception:
                db.session.rollback()
    except Exception as e:
        print(f"Error during database initialization: {str(e)}")

if __name__ == '__main__':
    app.run(debug=True, port=8888)
